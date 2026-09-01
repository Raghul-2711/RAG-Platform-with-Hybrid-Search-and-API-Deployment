"""
Loader for Microsoft PowerPoint PPTX documents using python-pptx.
"""

import os
from typing import List

from pptx import Presentation

from .base import DocumentLoader, LoadedDocument


class PptxLoader(DocumentLoader):
    """Loads PPTX files slide-by-slide."""

    SUPPORTED_EXTENSIONS = {".pptx"}

    def can_load(self, file_path: str) -> bool:
        """Return True if the file is a PPTX presentation."""
        extension = os.path.splitext(file_path)[1].lower()
        return extension in self.SUPPORTED_EXTENSIONS

    def load(self, file_path: str) -> List[LoadedDocument]:
        """Extract text from each PPTX slide as a normalized document."""
        if not self.can_load(file_path):
            raise ValueError(f"Unsupported PPTX file: {file_path}")

        filename = os.path.basename(file_path)
        presentation = Presentation(file_path)

        documents = []
        total_slides = len(presentation.slides)

        for slide_number, slide in enumerate(presentation.slides, start=1):
            slide_text = []

            for shape in slide.shapes:
                if not hasattr(shape, "text"):
                    continue

                text = shape.text.strip()
                if text:
                    slide_text.append(text)

            text = "\n".join(slide_text).strip()

            if not text:
                continue

            documents.append(
                LoadedDocument(
                    source=filename,
                    text=text,
                    metadata={
                        "document_name": filename,
                        "file_type": "pptx",
                        "slide_number": slide_number,
                        "total_slides": total_slides,
                    },
                )
            )

        return documents