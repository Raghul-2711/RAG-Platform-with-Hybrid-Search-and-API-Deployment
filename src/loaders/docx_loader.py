"""
Loader for Microsoft Word DOCX documents using python-docx.
"""

import os
from typing import List

from docx import Document

from .base import DocumentLoader, LoadedDocument


class DocxLoader(DocumentLoader):
    """Loads DOCX files and extracts paragraphs and tables."""

    SUPPORTED_EXTENSIONS = {".docx"}

    def can_load(self, file_path: str) -> bool:
        """Return True if the file is a DOCX document."""
        extension = os.path.splitext(file_path)[1].lower()
        return extension in self.SUPPORTED_EXTENSIONS

    def load(self, file_path: str) -> List[LoadedDocument]:
        """Extract text from a DOCX file as a normalized document."""
        if not self.can_load(file_path):
            raise ValueError(f"Unsupported DOCX file: {file_path}")

        filename = os.path.basename(file_path)
        document = Document(file_path)

        sections = []

        # Extract paragraphs.
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                sections.append(text)

        # Extract tables.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                row_text = " | ".join(cell for cell in cells if cell)

                if row_text:
                    sections.append(row_text)

        text = "\n".join(sections).strip()

        if not text:
            return []

        return [
            LoadedDocument(
                source=filename,
                text=text,
                metadata={
                    "document_name": filename,
                    "file_type": "docx",
                },
            )
        ]